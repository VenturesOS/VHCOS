"""Import the offline Branch & Recruiter Revenue tracker into `placement_ledger`.

The Excel (527 placements, FY 26-27) is the authoritative record of revenue
booked before the platform tracked joinings. Every target / dashboard number
is derived from this ledger + revenue booked on new joinings, which replaces
the hand-typed "already achieved" openings.

Mapping rules confirmed by the client (2026-09-20):
  • Spelling variants are one person (Karambir=Karamvir, Abhey=Abhay, ...).
  • Priyanka (hr64) and Priyanka yadav (hr9) are two different people.
  • People who left keep their revenue inside the team they worked for.
  • "Ajit / Avinash" is shared credit — split 50-50.
  • Blank recruiter / stray rows sit in the branch total as Unassigned.

Run:  python3 -m scripts.import_placement_ledger /path/to/file.xlsx
Idempotent — it clears rows of the same `source` first.
"""
import asyncio
import os
import re
import sys
import uuid
from collections import Counter
from datetime import date, timedelta

import openpyxl
from dotenv import load_dotenv
from motor.motor_asyncio import AsyncIOMotorClient

load_dotenv(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), ".env"))

SOURCE = "excel_tracker_2026_09"
COLL = "placement_ledger"

BRANCH_TEAM_NAME = {
    "Bangalore": "Bengaluru team",
    "Delhi": "Delhi Team",
    "Faridabad": "Faridabad team",
    "Gurgaon": "Gurgaon Team",
    "Hyderabad": "Krishna Team",
}

# Excel recruiter (normalised) → login email. `None` = no login, revenue stays
# in the branch under Ex-employee / Unassigned.
RECRUITER_EMAIL = {
    # Delhi
    "prabhatpanchal": "prabhat@vhc.in",
    "bhumika": "bhumika@vhc.in",
    "yamini": "yamini@vhc.in",
    "navyajain": "hr24@vhc.in",
    "stuti": "stuti@vhc.in",
    "kavita": "hr23@vhc.in",
    "srishti": "hr2@vhc.in",
    "srishtigupta": "hr2@vhc.in",
    "simran": "_deact_20260617t082002_simran@vhc.in",
    "sneha": "hr5@vhc.in",
    "maneet": "maneet@vhc.in",
    "khushimakkar": "hr21@vhc.in",
    "manas": "hr25@vhc.in",
    "shubham": None,
    # Gurgaon
    "jatin": "jatin@vhc.in",
    "jatinyadav": "jatin@vhc.in",
    "ajityadav": "ajit@vhc.in",
    "robin": "robin@vhc.in",
    "robinyadav": "robin@vhc.in",
    "priyanka": "hr64@vhc.in",
    "priyanaka": "hr64@vhc.in",
    "priyankayadav": "hr9@vhc.in",
    "saurabh": "hr10@vhc.in",
    "gaurav": "gaurav@vhc.in",
    "karamvir": "hr63@vhc.in",
    "karambir": "hr63@vhc.in",
    "abhay": "hr7@vhc.in",
    "abhey": "hr7@vhc.in",
    "rohit": "rohit@vhc.in",
    "rohityadav": "rohit@vhc.in",
    "janvi": "hr8@vhc.in",
    "janvisharma": "hr8@vhc.in",
    "ronak": "hr14@vhc.in",
    "sudha": "hr11@vhc.in",
    "diya": "hr6@vhc.in",
    "sachin": "hr12@vhc.in",
    "yash": "hr13@vhc.in",
    "jitender": None,
    "nikita": None,
    # Faridabad
    "indu": "hr53@vhc.in",
    "indusagar": "hr53@vhc.in",
    "madhuri": "hr58@vhc.in",
    "bikash": "bikash@vhc.in",
    "bikashdas": "bikash@vhc.in",
    "shivani": "hr57@vhc.in",
    "shivanidinkar": "hr57@vhc.in",
    "neha": "hr50@vhc.in",
    "nehamatpal": "hr50@vhc.in",
    "mukta": "hr56@vhc.in",
    "muktakumari": "hr56@vhc.in",
    "deepika": "hr52@vhc.in",
    "deepikatalwar": "hr52@vhc.in",
    "lalit": "_deact_20260519t110017_hr55@vhc.in",
    "lalitkumar": "_deact_20260519t110017_hr55@vhc.in",
    "nidhisingh": None,
    "parul": None,
    "parulmangla": None,
    # Bangalore
    "navya": "hr35@vhc.in",          # Bangalore Navya = Navya CN (see BRANCH_OVERRIDE)
    "rakshita": "hr37@vhc.in",
    "rakshitha": "hr37@vhc.in",
    "jayadev": "jayadev@vhc.in",
    "shwetha": "hr27@vhc.in",
    "avinash": "avinash@vhc.in",
    "vedha": "hr28@vhc.in",
    "visali": "hr36@vhc.in",
    "vishali": "hr36@vhc.in",
    "lalithasri": "hr29@vhc.in",
    "lalitha": "hr29@vhc.in",
    "haritha": "hr26@vhc.in",
    # Hyderabad
    "krishna": "krishna@vhc.in",
    "madhavi": "hr81@vhc.in",
    "madhavimancharia": "hr81@vhc.in",
}

# Same first name, different person depending on the branch.
BRANCH_OVERRIDE = {
    ("Delhi", "navya"): "hr24@vhc.in",       # Navya jain
    ("Bangalore", "navya"): "hr35@vhc.in",   # Navya CN
}

# Shared credit — revenue and the placement count are split evenly.
SPLIT = {"ajitavinash": ["ajit@vhc.in", "avinash@vhc.in"]}

UNASSIGNED = {"unassignedblankrecruiter", "areasalesmanager"}

# People with no login — one canonical display name per person.
EX_EMPLOYEE_NAME = {
    "parul": "Parul Mangla", "parulmangla": "Parul Mangla",
    "shubham": "Shubham", "jitender": "Jitender",
    "nikita": "Nikita", "nidhisingh": "Nidhi Singh",
}

MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}

STATUSES = {"payment received": "Payment Received", "pp": "PP", "ip": "IP",
            "backout": "Backout", "credit note": "Credit Note"}


def norm(s) -> str:
    return re.sub(r"[^a-z]", "", str(s or "").lower())


def parse_doj(raw) -> str:
    """`09.09.26` · `25.05.2026` · `30.6.2026` · `13.5-2026` → ISO."""
    if raw is None:
        return ""
    if hasattr(raw, "year"):
        return f"{raw.year:04d}-{raw.month:02d}-{raw.day:02d}"
    # Cells the sheet stored as a real date come back as an Excel serial
    if isinstance(raw, (int, float)) and 20000 < float(raw) < 60000:
        d = date(1899, 12, 30) + timedelta(days=int(raw))
        return d.isoformat()
    parts = [p for p in re.split(r"[.\-/]", str(raw).strip()) if p]
    if len(parts) != 3:
        # `16th July - 2026`
        m = re.match(r"\s*(\d{1,2})\w*\s+([A-Za-z]{3})[a-z]*\s*[-,]?\s*(\d{4})", str(raw))
        if m and m.group(2).lower() in MONTHS:
            return f"{int(m.group(3)):04d}-{MONTHS[m.group(2).lower()]:02d}-{int(m.group(1)):02d}"
        return ""
    try:
        d, m, y = int(parts[0]), int(parts[1]), int(parts[2])
    except ValueError:
        return ""
    if y < 100:
        y += 2000
    if not (1 <= m <= 12 and 1 <= d <= 31):
        return ""
    return f"{y:04d}-{m:02d}-{d:02d}"


def money(raw) -> float:
    if raw is None:
        return 0.0
    if isinstance(raw, (int, float)):
        return float(raw)
    cleaned = re.sub(r"[^0-9.\-]", "", str(raw))
    try:
        return float(cleaned)
    except ValueError:
        return 0.0


async def main(path: str):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Sorted Data"]
    header = [c for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    col = {name: header.index(name) for name in header if name}

    client = AsyncIOMotorClient(os.environ["MONGO_URL"])
    db = client[os.environ["DB_NAME"]]

    users = {}
    async for u in db.users.find({}, {"_id": 0, "id": 1, "name": 1, "email": 1, "is_active": 1}):
        users[str(u.get("email") or "").lower()] = u
    teams = {}
    async for t in db.teams.find({}, {"_id": 0, "id": 1, "name": 1, "employer_id": 1}):
        teams[t["name"]] = t

    docs, unknown, years = [], Counter(), Counter()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[col["S.No"]] is None:
            continue
        branch = row[col["Branch (Standard)"]]
        team = teams.get(BRANCH_TEAM_NAME.get(branch, ""), {})
        label = str(row[col["Recruiter (Standard)"]] or "").strip()
        key = norm(label)
        doj = parse_doj(row[col["DOJ"]])
        years[doj[:4] or "?"] += 1
        status_raw = str(row[col["Payment Status (Standard)"]] or "").strip()
        status = STATUSES.get(status_raw.lower(), "Other / Review")

        base = {
            "s_no": row[col["S.No"]],
            "branch": branch,
            "team_id": team.get("id") or "",
            "account_manager": row[col["Account Manager (Standard)"]] or "",
            "organization": row[col["Organization"]] or "",
            "candidate_name": row[col["Candidate Name"]] or "",
            "designation": row[col["Designation"]] or "",
            "location": row[col["Location"]] or "",
            "doj": doj,
            "doj_raw": str(row[col["DOJ"]] or ""),
            "offered_ctc": money(row[col["Offered CTC"]]),
            "invoice_no": str(row[col["Invoice No."]] or ""),
            "invoice_date_raw": str(row[col["Invoice Date"]] or ""),
            "payment_status": status,
            "payment_status_raw": status_raw,
            "review_note": row[col["Data Quality Note"]] or "",
            "source": SOURCE,
        }
        revenue = money(row[col["Revenue (Numeric)"]])

        if key in UNASSIGNED:
            targets = [(None, "Ex-employee / Unassigned", 1.0)]
        elif key in SPLIT:
            share = 1.0 / len(SPLIT[key])
            targets = [(e, None, share) for e in SPLIT[key]]
        else:
            email = BRANCH_OVERRIDE.get((branch, key), RECRUITER_EMAIL.get(key, "__missing__"))
            if email == "__missing__":
                unknown[f"{branch} | {label}"] += 1
                email = None
            targets = [(email, EX_EMPLOYEE_NAME.get(key, label) if email is None else None, 1.0)]

        for email, fallback_label, share in targets:
            u = users.get((email or "").lower(), {})
            docs.append({
                **base,
                "id": str(uuid.uuid4()),
                "recruiter_id": u.get("id") or "",
                "recruiter_name": u.get("name") or fallback_label or label,
                "recruiter_label": label,
                "is_ex_employee": bool(u) and (u.get("is_active") is False
                                              or str(u.get("email") or "").startswith("_deact_")),
                "unassigned": not u,
                "revenue": round(revenue * share, 2),
                "placement_credit": share,
                "shared_credit": share != 1.0,
            })

    removed = (await db[COLL].delete_many({"source": SOURCE})).deleted_count
    if docs:
        await db[COLL].insert_many(docs)
    await db[COLL].create_index([("doj", 1)])
    await db[COLL].create_index([("recruiter_id", 1), ("doj", 1)])
    await db[COLL].create_index([("team_id", 1), ("doj", 1)])

    print(f"removed={removed} inserted={len(docs)} placements={sum(d['placement_credit'] for d in docs)}")
    print("DOJ years:", dict(years))
    print("gross:", round(sum(d["revenue"] for d in docs), 2))
    if unknown:
        print("!! no mapping (kept as Unassigned):")
        for k, v in unknown.most_common():
            print("   ", k, v)


if __name__ == "__main__":
    asyncio.run(main(sys.argv[1] if len(sys.argv) > 1 else "/tmp/rev.xlsx"))
