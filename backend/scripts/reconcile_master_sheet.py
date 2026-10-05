"""Cross-check `placement_ledger` against the branch master sheet.

The master sheet has one tab per branch (Bangalore / Gurgaon / Delhi /
Faridabad / Hyderabad) with the same shape:

    S.No · Branch · Account Manager · Recruiter · Organization · Candidate ·
    Designation · [Department|Rate] · Location · DOJ · Offered CTC ·
    Invoice No. · Billing Amount · Invoice Date · Payment Status · [remarks]

Run read-only first:
    python3 -m scripts.reconcile_master_sheet /tmp/master.xlsx
Write the differences back:
    python3 -m scripts.reconcile_master_sheet /tmp/master.xlsx --apply
"""
import asyncio
import os
import re
import sys
import uuid
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta, timezone

import openpyxl
from dotenv import load_dotenv
from motor.motor_asyncio import AsyncIOMotorClient

load_dotenv(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), ".env"))

COLL = "placement_ledger"
SOURCE = "excel_tracker_2026_09"          # rows the first import created
MASTER_SOURCE = "master_sheet_2026_09_30"  # rows this run adds

# User decisions, 2026-09-30:
#  * a sheet row whose candidate is already a `joined` application stays
#    platform-owned, so Accounts can still raise its invoice in-app;
#  * RPO monthly-billing rows carry no DOJ — the invoice date places them;
#  * three unusable dates were confirmed by hand;
#  * two ledger rows are renamed duplicates of sheet rows and get voided.
DOJ_OVERRIDE = {
    "charudattaudayamangi": "2026-02-01",
    "vishatiwari": "2026-02-01",
    "saravanans": "2026-08-31",
}
VOID_AS_DUPLICATE = [
    ("saransh", "Superseded by 'Saransh Bansode' in the master sheet"),
    ("rahul", "Superseded by 'Faizal Raza' (same invoice VHC/26-27/26) in the master sheet"),
]

SHEET_BRANCH = {
    "Avinash-banglore": "Bangalore",
    "Ajit_GGN": "Gurgaon",
    "Delhi": "Delhi",
    "BIKAS DAS": "Faridabad",
    "KRISHNA_HYD": "Hyderabad",
}

BRANCH_TEAM_NAME = {
    "Bangalore": "Bengaluru team",
    "Delhi": "Delhi Team",
    "Faridabad": "Faridabad team",
    "Gurgaon": "Gurgaon Team",
    "Hyderabad": "Krishna Team",
}

STATUSES = {
    "payment received": "Payment Received",
    "received": "Payment Received",
    "payment pending": "PP",
    "pp": "PP",
    "invoice pending": "IP",
    "ip": "IP",
    "backout": "Backout",
    "credit note": "Credit Note",
}

MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}

HEADER_ALIASES = {
    "s.no": "s_no", "sr. no.": "s_no",
    "branch": "branch",
    "account manager": "account_manager",
    "recruiter name": "recruiter", "name of the recruiter": "recruiter",
    "organization": "organization",
    "candidate name": "candidate_name", "name of candidate": "candidate_name",
    "designation": "designation",
    "location": "location",
    "doj": "doj",
    "offered ctc": "offered_ctc", "offered amount": "offered_ctc",
    "invoice no.": "invoice_no", "invoice no": "invoice_no",
    "billing amount": "billing_amount",
    "invoice date": "invoice_date",
    "payment status": "payment_status",
    "payment received (yes/no)": "payment_status",
    "remarks": "remarks",
}


def norm(s) -> str:
    return re.sub(r"[^a-z]", "", str(s or "").lower())


def parse_doj(raw) -> str:
    if raw is None:
        return ""
    if hasattr(raw, "year"):
        return f"{raw.year:04d}-{raw.month:02d}-{raw.day:02d}"
    if isinstance(raw, (int, float)) and 20000 < float(raw) < 60000:
        return (date(1899, 12, 30) + timedelta(days=int(raw))).isoformat()
    text = str(raw).strip()
    m = re.match(r"\s*(\d{1,2})\w*\s*[-,/ ]\s*([A-Za-z]{3})[a-z]*\s*[-,/ ]?\s*(\d{2,4})", text)
    if m and m.group(2).lower() in MONTHS:
        y = int(m.group(3))
        return f"{y + 2000 if y < 100 else y:04d}-{MONTHS[m.group(2).lower()]:02d}-{int(m.group(1)):02d}"
    parts = [p for p in re.split(r"[.\-/ ]", text) if p]
    if len(parts) == 3:
        try:
            d, mth, y = int(parts[0]), int(parts[1]), int(parts[2])
        except ValueError:
            return ""
        if y < 100:
            y += 2000
        if 1 <= mth <= 12 and 1 <= d <= 31:
            return f"{y:04d}-{mth:02d}-{d:02d}"
    return ""


def money(raw) -> float:
    if raw is None:
        return 0.0
    if isinstance(raw, (int, float)):
        return round(float(raw), 2)
    cleaned = re.sub(r"[^0-9.\-]", "", str(raw).replace(",", ""))
    try:
        return round(float(cleaned), 2)
    except ValueError:
        return 0.0


def read_sheet(path: str) -> list:
    wb = openpyxl.load_workbook(path, data_only=True)
    rows = []
    for ws in wb.worksheets:
        branch = SHEET_BRANCH.get(ws.title)
        if not branch:
            print(f"  ! skipping unknown tab {ws.title!r}")
            continue
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        col = {}
        for i, h in enumerate(header):
            key = HEADER_ALIASES.get(str(h or "").strip().lower().rstrip(' '))
            if key and key not in col:
                col[key] = i
        missing = {"candidate_name", "organization", "billing_amount"} - set(col)
        if missing:
            print(f"  ! {ws.title}: missing columns {missing} — skipped")
            continue
        for raw in ws.iter_rows(min_row=2, values_only=True):
            def cell(name):
                i = col.get(name)
                return raw[i] if i is not None and i < len(raw) else None
            name = str(cell("candidate_name") or "").strip()
            if not name:
                continue
            status_raw = str(cell("payment_status") or "").strip()
            doj = parse_doj(cell("doj")) or DOJ_OVERRIDE.get(norm(name), "")
            if not doj and "rpo" in name.lower():
                doj = parse_doj(cell("invoice_date"))
            rows.append({
                "tab": ws.title,
                "s_no": cell("s_no"),
                "branch": branch,
                "account_manager": str(cell("account_manager") or "").strip(),
                "recruiter_label": str(cell("recruiter") or "").strip(),
                "organization": str(cell("organization") or "").strip(),
                "candidate_name": name,
                "designation": str(cell("designation") or "").strip(),
                "location": str(cell("location") or "").strip(),
                "doj": doj,
                "doj_raw": str(cell("doj") or ""),
                "offered_ctc": money(cell("offered_ctc")),
                "invoice_no": str(cell("invoice_no") or "").strip(),
                "billing_amount": money(cell("billing_amount")),
                "invoice_date_raw": str(cell("invoice_date") or ""),
                "payment_status": STATUSES.get(status_raw.lower(), "Other / Review" if status_raw else "IP"),
                "payment_status_raw": status_raw,
                "remarks": str(cell("remarks") or "").strip(),
            })
    return rows


async def load_users_teams(db):
    users_by_norm = {}
    async for u in db.users.find({}, {"_id": 0, "id": 1, "name": 1, "email": 1, "is_active": 1}):
        users_by_norm.setdefault(norm(u.get("name")), u)
    teams = {}
    async for t in db.teams.find({}, {"_id": 0, "id": 1, "name": 1}):
        teams[t["name"]] = t
    return users_by_norm, teams


async def main(path: str, apply: bool):
    sheet_rows = read_sheet(path)
    print(f"sheet rows: {len(sheet_rows)}")
    print("by branch:", dict(Counter(r["branch"] for r in sheet_rows)))
    print("by status:", dict(Counter(r["payment_status"] for r in sheet_rows)))
    print("gross billing in sheet:", round(sum(r["billing_amount"] for r in sheet_rows), 2))
    no_doj = [r for r in sheet_rows if not r["doj"]]
    if no_doj:
        print(f"! {len(no_doj)} rows with an unparseable DOJ:",
              [(r["candidate_name"], r["doj_raw"]) for r in no_doj][:10])

    client = AsyncIOMotorClient(os.environ["MONGO_URL"])
    db = client[os.environ["DB_NAME"]]

    ledger = await db[COLL].find({"void": {"$ne": True}}, {"_id": 0}).to_list(20000)
    print(f"\nledger rows: {len(ledger)} | gross {round(sum(float(r.get('revenue') or 0) for r in ledger), 2)}")

    by_key = defaultdict(list)
    for r in ledger:
        by_key[(norm(r.get("candidate_name")), norm(r.get("organization")))].append(r)
    by_name = defaultdict(list)
    for r in ledger:
        by_name[norm(r.get("candidate_name"))].append(r)

    matched, mismatches, missing = [], [], []
    for s in sheet_rows:
        hits = by_key.get((norm(s["candidate_name"]), norm(s["organization"])))
        if not hits:
            hits = by_name.get(norm(s["candidate_name"]))
        if not hits:
            missing.append(s)
            continue
        matched.append((s, hits))
        booked = round(sum(float(h.get("revenue") or 0) for h in hits), 2)
        diff = {}
        if abs(booked - s["billing_amount"]) > 1:
            diff["billing"] = (booked, s["billing_amount"])
        db_status = hits[0].get("payment_status") or ""
        if db_status != s["payment_status"]:
            diff["status"] = (db_status, s["payment_status"])
        db_inv = str(hits[0].get("invoice_no") or "").strip()
        if s["invoice_no"] and db_inv != s["invoice_no"]:
            diff["invoice"] = (db_inv, s["invoice_no"])
        db_doj = str(hits[0].get("doj") or "")
        if s["doj"] and db_doj != s["doj"]:
            diff["doj"] = (db_doj, s["doj"])
        if diff:
            mismatches.append((s, hits, diff))

    print(f"\nmatched: {len(matched)} | mismatched fields: {len(mismatches)} | only in sheet: {len(missing)}")
    print("\n--- ONLY IN SHEET (would be inserted) ---")
    for s in missing[:60]:
        print(f"  {s['branch']:10} {s['doj'] or s['doj_raw']:12} {s['candidate_name'][:28]:30} "
              f"{s['organization'][:22]:24} {s['billing_amount']:>12,.0f} {s['payment_status']}")
    if len(missing) > 60:
        print(f"  ... and {len(missing) - 60} more")

    print("\n--- FIELD MISMATCHES (db → sheet) ---")
    kinds = Counter()
    for s, hits, diff in mismatches[:80]:
        kinds.update(diff.keys())
        parts = " | ".join(f"{k}: {a!r} → {b!r}" for k, (a, b) in diff.items())
        print(f"  {s['branch']:10} {s['candidate_name'][:26]:28} {parts}")
    for s, hits, diff in mismatches[80:]:
        kinds.update(diff.keys())
    if len(mismatches) > 80:
        print(f"  ... and {len(mismatches) - 80} more")
    print("mismatch kinds:", dict(kinds))

    sheet_names = {norm(r["candidate_name"]) for r in sheet_rows}
    extra = [r for r in ledger if norm(r.get("candidate_name")) not in sheet_names]
    print(f"\n--- ONLY IN THE LEDGER (left untouched) --- {len(extra)} rows, "
          f"gross {round(sum(float(r.get('revenue') or 0) for r in extra), 2)}")
    print("   by branch:", dict(Counter(r.get("branch") for r in extra)))

    if not apply:
        print("\nread-only run — nothing written. Re-run with --apply to write.")
        return

    users_by_norm, teams = await load_users_teams(db)
    now = datetime.now(timezone.utc).isoformat()
    updated = 0
    for s, hits, diff in mismatches:
        target = hits[0]
        patch = {
            "payment_status": s["payment_status"],
            "payment_status_raw": s["payment_status_raw"],
            "invoice_no": s["invoice_no"] or target.get("invoice_no", ""),
            "invoice_date_raw": s["invoice_date_raw"] or target.get("invoice_date_raw", ""),
            "offered_ctc": s["offered_ctc"] or target.get("offered_ctc", 0),
            "master_synced_at": now,
        }
        if s["doj"]:
            patch["doj"] = s["doj"]
            patch["doj_raw"] = s["doj_raw"]
        if "billing" in diff and len(hits) == 1:
            # Split rows keep their share; only a single owner gets the reset.
            patch["revenue"] = s["billing_amount"]
        await db[COLL].update_one({"id": target["id"]}, {"$set": patch})
        updated += 1

    joined_names = set()
    async for a in db.applications.find({"stage": "joined"}, {"_id": 0, "candidate_name": 1}):
        joined_names.add(norm(a.get("candidate_name")))

    inserted, skipped_platform = [], []
    for s in missing:
        if norm(s["candidate_name"]) in joined_names:
            skipped_platform.append(s)
            continue
        team = teams.get(BRANCH_TEAM_NAME.get(s["branch"], ""), {})
        u = users_by_norm.get(norm(s["recruiter_label"]), {})
        inserted.append({
            "id": str(uuid.uuid4()),
            "s_no": s["s_no"],
            "branch": s["branch"],
            "team_id": team.get("id") or "",
            "account_manager": s["account_manager"],
            "organization": s["organization"],
            "candidate_name": s["candidate_name"],
            "designation": s["designation"],
            "location": s["location"],
            "doj": s["doj"],
            "doj_raw": s["doj_raw"],
            "offered_ctc": s["offered_ctc"],
            "invoice_no": s["invoice_no"],
            "invoice_date_raw": s["invoice_date_raw"],
            "payment_status": s["payment_status"],
            "payment_status_raw": s["payment_status_raw"],
            "review_note": s["remarks"],
            "recruiter_id": u.get("id") or "",
            "recruiter_name": u.get("name") or s["recruiter_label"] or "Ex-employee / Unassigned",
            "recruiter_label": s["recruiter_label"],
            "is_ex_employee": bool(u) and (u.get("is_active") is False
                                           or str(u.get("email") or "").startswith("_deact_")),
            "unassigned": not u,
            "revenue": s["billing_amount"],
            "placement_credit": 1.0,
            "shared_credit": False,
            "source": MASTER_SOURCE,
            "master_synced_at": now,
        })
    if inserted:
        await db[COLL].insert_many(inserted)
    voided = 0
    for name_norm, note in VOID_AS_DUPLICATE:
        res = await db[COLL].update_many(
            {"candidate_name": {"$regex": f"^{name_norm}$", "$options": "i"}, "void": {"$ne": True}},
            {"$set": {"void": True, "void_reason": note, "master_synced_at": now}},
        )
        voided += res.modified_count

    print(f"\nAPPLIED — updated {updated} rows, inserted {len(inserted)} rows, "
          f"voided {voided} duplicates")
    print(f"left platform-owned (already a joined application): {len(skipped_platform)} rows "
          f"₹{round(sum(r['billing_amount'] for r in skipped_platform), 2):,}")
    unmapped = Counter(r["recruiter_label"] for r in inserted if not r["recruiter_id"])
    if unmapped:
        print("inserted without a recruiter login (sit in the branch as Unassigned):",
              dict(unmapped))


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    asyncio.run(main(args[0] if args else "/tmp/master.xlsx", "--apply" in sys.argv))
