"""
submission_note.py — client-ready candidate submission note (DOCX)
==================================================================
One-page VHC-branded note generated from the candidate record + Asha's
verified answers. python-docx only (add to requirements). The recruiter
downloads it from the session drawer and pastes into the client tracker
or attaches to the submission mail.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any, Dict

NAVY = "0F2A4A"
GOLD = "C9A54A"


def build_submission_note(candidate: Dict[str, Any], session: Dict[str, Any],
                          job: Dict[str, Any], out_dir: str = "/tmp") -> str:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    v = session.get("verified") or {}
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.6)
        s.left_margin = s.right_margin = Inches(0.7)

    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run("VENTURES HRD CENTRE")
    r.font.size, r.font.bold = Pt(16), True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = sub.add_run("Candidate Submission Note — screened & verified by Asha")
    r2.font.size = Pt(10)
    r2.font.color.rgb = RGBColor.from_string(GOLD)

    doc.add_paragraph()
    t = doc.add_paragraph()
    rt = t.add_run(f"{candidate.get('name', 'Candidate')}  —  "
                   f"{job.get('title') or session.get('mandate_title') or ''}")
    rt.font.size, rt.font.bold = Pt(13), True
    rt.font.color.rgb = RGBColor.from_string(NAVY)

    def _fmt(x, suffix=""):
        return f"{x}{suffix}" if x not in (None, "", []) else "—"

    rows = [
        ("Mobile", _fmt(candidate.get("phone"))),
        ("Email", _fmt(candidate.get("email"))),
        ("Current company", _fmt(candidate.get("current_employer") or candidate.get("current_company"))),
        ("Designation", _fmt(candidate.get("current_designation") or candidate.get("designation"))),
        ("Total experience", _fmt(v.get("total_experience_years") or candidate.get("total_experience_years"), " years")),
        ("Current CTC (verified)", _fmt(v.get("current_lpa"), " LPA")),
        ("Expected CTC (verified)", _fmt(v.get("expected_lpa"), " LPA")),
        ("Notice period (verified)", _fmt(v.get("notice_days"), " days")),
        ("Location / relocation", _fmt({"there": "At job location", "willing": "Willing to relocate",
                                        "not_willing": "Not willing to relocate"}.get(v.get("relocation_status"),
                                                                                      candidate.get("current_location")))),
        ("Skills verified in screening", ", ".join(v.get("verified_skills") or []) or "—"),
        ("Match score", _fmt(session.get("score"), "%")),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for label, value in rows:
        cells = table.add_row().cells
        p0 = cells[0].paragraphs[0].add_run(label)
        p0.font.bold, p0.font.size = True, Pt(9.5)
        p1 = cells[1].paragraphs[0].add_run(str(value))
        p1.font.size = Pt(9.5)

    doc.add_paragraph()
    note = doc.add_paragraph()
    nr = note.add_run(
        "Verified fields were confirmed directly with the candidate in a structured "
        "WhatsApp screening conversation (transcript on record with VHC). "
        f"Screened on {str(session.get('completed_at') or '')[:10]}.")
    nr.font.size = Pt(8.5)
    nr.font.color.rgb = RGBColor.from_string("54677E")

    safe = "".join(c for c in (candidate.get("name") or "candidate") if c.isalnum() or c in " _-").strip().replace(" ", "_")
    path = str(Path(out_dir) / f"VHC_Submission_{safe}_{session['id'][:8]}.docx")
    doc.save(path)
    return path
